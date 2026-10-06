#!/usr/bin/env python3
"""Codification automatique d'articles consommables (Excel).
Principe : la recherche par mots (sur le Referentiel + lignes validees) decide seule quand c'est
evident ; l'IA locale (Ollama) n'intervient qu'en dernier recours, avec un contexte minimal."""
import csv
import json
import math
import os
import re
import sys
import time
import unicodedata
from collections import Counter
from copy import copy
from datetime import datetime
from pathlib import Path

import requests
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.formula.translate import Translator
from openpyxl.styles import PatternFill

# ----------------------------- Constantes ---------------------------------
MODELE = "qwen2.5:3b"   # plus rapide mais moins precis : "qwen2.5:1.5b"
URL_OLLAMA = "http://localhost:11434"
BASE = Path(__file__).resolve().parent
DOSSIER_INPUT, DOSSIER_OUTPUT = BASE / "input", BASE / "output"
JAUNE = "FFFF00"        # marque "a traiter" posee par l'utilisateur sur la colonne A
LONGUEUR_MAX = 30
NUM_CTX, NUM_PREDICT = 2048, 40               # contexte minimal = RAM / CPU legers
NUM_THREAD = max(2, (os.cpu_count() or 4) // 2)  # coeurs physiques (hyper-threading inutile)
SEUIL_DIRECT, RATIO_DIRECT = 6.0, 1.6           # choix sans IA si score >= seuil et >= ratio x 2e
NB_CANDIDATES = 6
LIGNES_REFERENTIEL = 500  # les plages VLOOKUP vers le Referentiel sont elargies jusque-la
# role -> (regex cherchee dans l'en-tete de la ligne 1, colonne par defaut)
ROLES = {"K": ("code famille", 11), "M": ("code sous", 13), "N": ("^detail", 14),
         "statut": ("statut", 16), "collision": ("collision", 17),
         "nim": ("nimroda", 2), "four": ("designation fournisseur", 3), "ref": ("^ref", 4)}
PALETTE = {"jaune": "FFFF00", "rouge": "FF0000", "rouge clair": "F8D7DA", "bleu clair": "E8F0FE",
           "bleu fonce": "1F4E78", "orange clair": "FFD8A8", "jaune clair": "FFF3CD",
           "vert clair": "C6EFCE", "gris": "D9D9D9", "blanc": "FFFFFF"}
ROUGE, VERT, GRIS, FIN = "\x1b[31m", "\x1b[32m", "\x1b[90m", "\x1b[0m"


# ----------------------------- Utilitaires --------------------------------
def txt(v):
    return "" if v is None else str(v).strip()


def norm(t):
    """Minuscules sans accents."""
    s = unicodedata.normalize("NFKD", txt(t).lower())
    return "".join(c for c in s if not unicodedata.combining(c))


def mots(t):
    """Mots (lettres, >= 3) au singulier approximatif."""
    return {w[:-1] if len(w) > 4 and w[-1] in "sx" else w for w in re.findall(r"[a-z]{3,}", norm(t))}


def nettoyer(detail):
    """Normalise un detail : 10*330 / 10X330 -> 10.330, 3,7 -> 3.7, majuscules, A-Z 0-9 . - +"""
    s = re.sub(r"(?<=\d)[*xX](?=\d)", ".", norm(detail))
    s = re.sub(r"(?<=\d),(?=\d)", ".", s)
    return re.sub(r"[^A-Z0-9.+\-]", "", s.upper())


def couleur(cell):
    """Couleur de fond RGB (6 hexa) ou None."""
    f = cell.fill
    if f and f.fill_type == "solid" and f.fgColor.type == "rgb":
        return str(f.fgColor.rgb)[-6:].upper()
    return None


def nom_couleur(rgb):
    p = [int(rgb[i:i + 2], 16) for i in (0, 2, 4)]
    return min(PALETTE, key=lambda n: sum((int(PALETTE[n][i:i + 2], 16) - p[j]) ** 2 for j, i in enumerate((0, 2, 4))))


def pastille(rgb):
    return f"\x1b[48;2;{int(rgb[0:2], 16)};{int(rgb[2:4], 16)};{int(rgb[4:6], 16)}m  {FIN}"


def verifier_ollama():
    """Arrete le programme si Ollama / le modele sont absents, puis charge le modele en memoire."""
    try:
        noms = [m["name"] for m in requests.get(f"{URL_OLLAMA}/api/tags", timeout=5).json()["models"]]
    except Exception:
        sys.exit("ERREUR : Ollama ne repond pas. Lance Ollama puis relance ce programme.")
    if not any(n == MODELE or n.startswith(MODELE + ":") for n in noms):
        sys.exit(f"ERREUR : modele absent. Tape dans un terminal :  ollama pull {MODELE}")
    print("Chargement du modele en memoire...", flush=True)
    requests.post(f"{URL_OLLAMA}/api/chat", timeout=600, json={
        "model": MODELE, "stream": False, "keep_alive": "30m", "options": {"num_predict": 1, "num_ctx": NUM_CTX},
        "messages": [{"role": "user", "content": "ok"}]})


def appeler_ia(fixe, variable, titre, choix=None):
    """Appel Ollama independant (aucun historique). Partie fixe d'abord : le cache Ollama la reutilise.
    Si `choix` est donne, la reponse est contrainte a {"cle": <une valeur de choix>}."""
    print(f"     {GRIS}[{titre}] ENVOI ({len(fixe) + len(variable)} car.){FIN} " + variable.replace("\n", " / "), flush=True)
    debut = time.time()
    forme = {"type": "object", "properties": {"cle": {"type": "string", "enum": choix}}, "required": ["cle"]} if choix else "json"
    r = requests.post(f"{URL_OLLAMA}/api/chat", timeout=600, json={
        "model": MODELE, "stream": False, "format": forme, "keep_alive": "30m",
        "options": {"temperature": 0, "num_ctx": NUM_CTX, "num_predict": NUM_PREDICT, "num_thread": NUM_THREAD},
        "messages": [{"role": "user", "content": fixe + "\n\n" + variable}]})
    r.raise_for_status()
    brut = r.json()["message"]["content"]
    print(f"     {GRIS}[{titre}] REPONSE en {time.time() - debut:.1f}s :{FIN} {brut.strip()}", flush=True)
    try:
        rep = json.loads(brut)
        return rep if isinstance(rep, dict) else {}
    except ValueError:
        return {}


# ----------------------------- Lecture du fichier -------------------------
def lire_legende(wb):
    """Onglet Legende : cellule coloree vide + texte a droite. Renvoie {rgb: sens} et les couleurs par role."""
    sens = {}
    if "Legende" in wb.sheetnames:
        for row in wb["Legende"].iter_rows():
            for c in row:
                rgb = couleur(c)
                if rgb and not txt(c.value) and txt(c.offset(0, 1).value):
                    sens[rgb] = txt(c.offset(0, 1).value)
    roles = {"edit": ("modifier", "a la main"), "collision": ("collision",),
             "traiter": ("a traiter", "a codifier", "a completer")}
    return sens, {r: {k for k, v in sens.items() if any(m in norm(v) for m in ms)} for r, ms in roles.items()}


def lire_referentiel(ws):
    """(familles {code: libelle}, sous-familles {cle: libelle}) lus dans le fichier."""
    fam, sub = {}, {}
    for a, b, _, d, e in ws.iter_rows(min_row=2, max_col=5, values_only=True):
        if txt(a):
            fam[txt(a).upper()] = txt(b)
        if txt(d):
            sub[txt(d).upper()] = txt(e)
    return fam, sub


def trouver_colonnes(ws, couleurs_edit):
    """Colonnes par en-tete ; a defaut, les 3 colonnes designees 'a modifier' par la legende."""
    col = {}
    for role, (mot, defaut) in ROLES.items():
        col[role] = next((c for c in range(1, ws.max_column + 1) if re.search(mot, norm(ws.cell(1, c).value))), defaut)
    editables = [c for c in range(1, ws.max_column + 1)
                 if any(couleur(ws.cell(r, c)) in couleurs_edit for r in range(2, min(ws.max_row, 40) + 1))]
    if len(editables) == 3 and not all(norm(ws.cell(1, c).value) for c in editables):
        col["K"], col["M"], col["N"] = editables
    return col, editables


def ligne_info(ws, r, col):
    g = lambda c: txt(ws.cell(r, c).value)
    return {"ligne": r, "num": g(1), "nim": g(col["nim"]), "four": g(col["four"]), "ref": g(col["ref"]),
            "K": g(col["K"]).upper(), "M": g(col["M"]).upper(), "N": g(col["N"]).upper()}


def code_de(i):
    return f"{i['K']}-{i['M']}-{i['N']}" if i["K"] and i["M"] and i["N"] else ""


def texte_article(i):
    return i["nim"] + ("" if i["four"].upper() == "X" else " " + i["four"])


# ----------------------------- Recherche (sans IA) ------------------------
def construire_index(sub, modeles):
    """Profil de mots par sous-famille : libelle du Referentiel (poids 3) + designations validees (poids 1)."""
    prof = {k: {w: 3.0 for w in mots(lib)} for k, lib in sub.items()}
    for m in modeles:
        if f"{m['K']}-{m['M']}" in prof:
            for w in mots(texte_article(m)):
                prof[f"{m['K']}-{m['M']}"].setdefault(w, 1.0)
    df = Counter(w for p in prof.values() for w in p)
    return prof, df


def classer(i, index):
    """[(score, cle)] decroissant. Un mot de l'article correspond a un mot du profil s'il est identique
    ou si l'un est le debut de l'autre (>= 4 lettres, poids / 2) : 'tele' ~ 'televiseur'."""
    prof, df = index
    art, res = mots(texte_article(i)), []
    for k, p in prof.items():
        s = 0.0
        for w in art:
            s += max((poids * math.log(1 + len(prof) / df[pw]) * (1 if pw == w else 0.5) for pw, poids in p.items()
                      if pw == w or (min(len(w), len(pw)) >= 4 and (pw.startswith(w) or w.startswith(pw)))), default=0.0)
        res.append((s, k))
    return sorted(res, reverse=True)


def proches(i, candidats, n):
    """Les n lignes modeles les plus proches de l'article (jetons lettres + chiffres communs)."""
    jet = lambda x: set(re.findall(r"[a-z0-9]+", norm(texte_article(x))))
    a = jet(i)
    return sorted(candidats, key=lambda m: -len(a & jet(m)))[:n]


# ----------------------------- Etape 1 : sous-famille ---------------------
def choisir_sous_famille(i, classement, sub, modeles):
    """Renvoie (cle, sure, famille). Direct (sure) si evident ; sinon l'IA tranche entre quelques candidates
    (sure seulement si elle confirme la recherche) ; si elle ne sait pas, on garde la meilleure candidate de la recherche."""
    (s1, k1), s2 = classement[0], classement[1][0] if len(classement) > 1 else 0
    famille = k1.split("-")[0] if s1 > 0 else None
    if s1 >= SEUIL_DIRECT and s1 >= RATIO_DIRECT * s2:
        print(f"  Etape 1/2 : {VERT}{k1} (choix direct, score {s1:.1f} contre {s2:.1f}, 0 appel IA){FIN}", flush=True)
        return k1, True, famille
    cand = [k for s, k in classement[:NB_CANDIDATES]] if s1 > 0 else list(sub)
    print(f"  Etape 1/2 : ambigu (score {s1:.1f} contre {s2:.1f}) -> IA sur {len(cand)} candidates", flush=True)
    lignes = []
    for k in cand:
        ex = proches(i, [m for m in modeles if f"{m['K']}-{m['M']}" == k], 1)
        lignes.append(f"{k} : {sub[k]}" + (f" (ex: {ex[0]['nim'][:40]})" if ex else ""))
    fixe = 'Choisis la cle de sous-famille de l\'article. Reponds en JSON : {"cle": "XXX-YY"}.\nCles :\n' + "\n".join(lignes)
    var = f"Article : {texte_article(i)[:200]}"
    for essai in range(2):
        cle = txt(appeler_ia(fixe, var + ("\nreponds uniquement par une cle de la liste." if essai else ""),
                             f"etape 1, essai {essai + 1}").get("cle")).upper()
        if cle in sub:
            return cle, cle == k1, famille  # sure si l'IA et la recherche sont d'accord
    return (k1, False, famille) if s1 > 0 else (None, False, None)


# ----------------------------- Etape 2 : detail par regles ----------------
# Regles deduites de l'Excel (section 5) : chaque TYPE de detail est un format applique a des attributs lus
# dans les designations (fournisseur d'abord). Le type est fixe par sous-famille, jamais par ligne.
N = r"(\d+(?:[.,]\d+)?)"
KX = r"\s*[x*\u00d7]\s*"
TYPES = {  # type -> (format, description donnee a l'IA quand aucune ligne validee n'existe)
    "DIAM": (lambda a: (a["diam"] or a["dims"] and a["dims"][0]) and "D" + (a["diam"] or a["dims"][0]),
             "diametre seul : D6.5 (forets, brosses)"),
    "TAILLE": (lambda a: a["taille"] and "T" + a["taille"], "taille : T10 (gants, chaussures, vetements)"),
    "TIGE": (lambda a: a["dims"] and a["g"] and f"D{a['dims'][0]}.{a['dims'][1]}-G{a['g']}",
             "outil rotatif sur tige (roue, manchon) : D30.15-G80"),
    "DISQUE": (lambda a: a["diam"] and a["g"] and f"D{a['diam']}G{a['g']}", "disque : diametre et grain, D125G36"),
    "PLAT": (lambda a: a["dims"] and a["g"] and f"{a['dims'][0]}.{a['dims'][1]}G{a['g']}",
             "produit plat (bande, feuille, capuchon) : dim1.dim2 et grain, 23.28G400"),
    "GRAIN": (lambda a: a["g"] and "G" + a["g"], "grain seul : G80 (rouleaux, disques de poncage)"),
    "COTE": (lambda a: a["dims"] and f"D{a['dims'][0]}.{a['dims'][1]}", "support : cotes du consommable porte, D6.10"),
    "VOLUME": (lambda a: a["vol"] and a["vol"] + "L", "produit liquide : volume, 5L"),
    "LARGEUR": (lambda a: a["mm"] and "L" + a["mm"], "ruban, rouleau adhesif : L25 (largeur en mm)"),
    "LUMENS": (lambda a: a["lum"] and "L" + a["lum"], "lampe : L300 (lumens)"),
    "STD": (lambda a: "STD", "outil ou EPI unique sans variante : STD"),
}


def attributs(t):
    """Attributs lus dans UNE designation (deja normalisee par norm)."""
    m = lambda p: re.search(p, t)
    diam = next((_f(x.group(1)) for x in (m(r"(?:diametre|diam|dia)\.?\s*:?\s*" + N), m(r"\u00f8\s*" + N),
                                          m(r"\bd\.?\s*" + N + r"\s*mm"), m(N + r"\s*mm"), m(r"\bd\.?\s*" + N)) if x), None)
    dims = m(N + KX + N + "(?:" + KX + N + ")?") or m(r"\bd\s*" + N + r"\s*l\s*" + N) or m(r"\b(\d{1,3})\.(\d{1,3})\b(?!\s*mm)")
    g = m(r"\b(?:grain|gr|g|p)\.?\s*[:.]?\s*(?:a|ceramique|oxy\w*|alu\w*)?\s*" + N + r"\s*(\+)?")
    mot = "FIN" if m(r"\b(ultra |tres |very )?fin(e)?\b") else "MOY" if m(r"moyen|medium") else None
    ta = m(r"(?:taille|pointure)\s*[.:]?\s*" + N + r"(?:/(\d{1,2}(?:[.,]\d)?)(?!\d))?") or m(r"\bt\s*(\d{1,2})\b")
    vol, mm, lum = m(N + r"\s*(?:l|litres?)\b"), m(N + r"\s*mm"), m(N + r"\s*lumens?")
    return {"mm": mm and _f(mm.group(1)), "lum": lum and _f(lum.group(1)), "diam": diam, "dims": [_f(x) for x in dims.groups() if x] if dims else None,
            "g": mot or (_f(g.group(1)) + (g.group(2) or "") if g else None),
            "taille": ta and "".join(_f(x) for x in ta.groups() if x), "vol": vol and _f(vol.group(1))}


def _f(x):
    return x.replace(",", ".")


def extraire(i):
    """Attributs de l'article : designation fournisseur d'abord, NIMRODA pour completer ; '+' du grain si l'une le porte."""
    a = {}
    for t in (norm(i["four"]) if i["four"].upper() != "X" else "", norm(i["nim"])):
        for k, v in attributs(t).items():
            if v and not a.get(k):
                a[k] = v
        if re.search(r"\d\s*\+", t) and a.get("g", "").isdigit():
            a["g"] += "+"
    return {k: a.get(k) for k in ("diam", "dims", "g", "taille", "vol", "mm", "lum")}


def apprendre_types(lignes, seuil):
    """{cle: type} : pour chaque sous-famille, le type dont le format reproduit les codes des lignes
    (None si aucun ne les reproduit : codes propres au produit)."""
    par = {}
    for m in lignes:
        if code_de(m):
            par.setdefault(f"{m['K']}-{m['M']}", []).append(m)
    out = {}
    for cle, ex in par.items():
        taux = {t: sum(f(extraire(m)) == nettoyer(m["N"]) for m in ex) / len(ex) for t, (f, _) in TYPES.items()}
        t = max(taux, key=taux.get)
        out[cle] = t if taux[t] >= seuil else None
    return out


def type_par_ia(cle, sub, i, types):
    """Sous-famille sans aucun code connu : l'IA choisit le type par analogie avec les sous-familles deja apprises."""
    fixe = ("Quel type de detail convient a la sous-famille ? Reponds en JSON {\"cle\": \"TYPE\"}.\nTypes :\n"
            + "\n".join(f"{t} : {d}" for t, (_, d) in TYPES.items()) + "\n\nDeja etablis :\n"
            + "\n".join(f"- {sub[k]} -> {t}" for k, t in types.items() if t and k in sub))
    return txt(appeler_ia(fixe, f"Sous-famille : {sub[cle]}\nArticle : {texte_article(i)[:160]}", "type de detail", list(TYPES)).get("cle"))


def construire_detail(i, cle, typ, utilises):
    """Renvoie (detail, probleme) ; probleme = None si le code est valide et unique."""
    d = TYPES[typ][0](extraire(i)) if typ else None
    probleme = ("detail propre au produit" if not typ else "attribut introuvable" if not d else
                "code trop long" if len(f"{cle}-{d}") > LONGUEUR_MAX else "collision" if f"{cle}-{d}" in utilises else None)
    print(f"  Etape 2/2 : {VERT if not probleme else ROUGE}detail {d or '?'} (type {typ}){FIN}", flush=True)
    return d or None, probleme


# ----------------------------- Ecriture dans le classeur ------------------
def elargir_plages(ws):
    """Plages figees vers le Referentiel (ex. $A$2:$B$7) -> elargies, pour que tout ajout soit reconnu."""
    n = 0
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("=") and "Referentiel!" in c.value:
                new = re.sub(r"(Referentiel!\$[A-Z]+\$\d+:\$[A-Z]+\$)(\d+)",
                             lambda m: m.group(1) + str(max(int(m.group(2)), LIGNES_REFERENTIEL)), c.value)
                if new != c.value:
                    c.value, n = new, n + 1
    return n


def reparer_ligne(ws, r, modele, cols_formules):
    """Ligne X sans formules : recopie formules et style d'une ligne modele (retire le fond rouge)."""
    for c in cols_formules:
        if not txt(ws.cell(r, c).value).startswith("="):
            ws.cell(r, c).value = Translator(ws.cell(modele, c).value, origin=ws.cell(modele, c).coordinate) \
                .translate_formula(ws.cell(r, c).coordinate)
    for c in range(2, ws.max_column + 1):
        ws.cell(r, c)._style = copy(ws.cell(modele, c)._style)


def afficher_ligne(ws, r, sens):
    """Affiche la ligne entiere (valeurs de A a la derniere colonne) et ses couleurs."""
    print(f"\n>>> Ligne {r}/{ws.max_row}")
    champs, couleurs = [], {}
    for c in range(1, ws.max_column + 1):
        cell, v = ws.cell(r, c), txt(ws.cell(r, c).value)
        if v and not v.startswith("="):
            champs.append(f"{txt(ws.cell(1, c).value)[:22]}: {v[:80]}")
        if couleur(cell):
            couleurs.setdefault(couleur(cell), []).append(cell.column_letter)
    print("    " + " | ".join(champs))
    for rgb, lettres in couleurs.items():
        s = sens.get(rgb) or ("marque 'a traiter'" if rgb == JAUNE else "")
        print(f"    {pastille(rgb)} {nom_couleur(rgb)} (#{rgb}) colonnes {','.join(lettres)}" + (f" = {s}" if s else ""))


def traiter_fichier(chemin, horodatage, rapport):
    wb = load_workbook(chemin)  # sans data_only : on garde toutes les formules
    if "Codification" not in wb.sheetnames or "Referentiel" not in wb.sheetnames:
        print("  -> onglets 'Codification' / 'Referentiel' introuvables, fichier ignore.")
        return
    ws, (fam, sub) = wb["Codification"], lire_referentiel(wb["Referentiel"])
    sens, roles = lire_legende(wb)
    col, editables = trouver_colonnes(ws, roles["edit"])
    lettre = lambda c: ws.cell(1, c).column_letter
    print(f"  Legende : {len(sens)} couleurs comprises ; cellules a modifier = colonnes "
          f"{', '.join(lettre(c) for c in editables) or '(non trouvees)'} ; utilise : "
          f"{lettre(col['K'])}, {lettre(col['M'])}, {lettre(col['N'])} (+ Statut {lettre(col['statut'])}, Collision {lettre(col['collision'])})")
    print(f"  Referentiel : {len(fam)} familles, {len(sub)} sous-familles ; {elargir_plages(ws)} formule(s) VLOOKUP elargie(s).")
    marques = {JAUNE} | roles["traiter"]
    infos = [ligne_info(ws, r, col) for r in range(2, ws.max_row + 1) if txt(ws.cell(r, col["nim"]).value)]
    a_traiter = [i for i in infos if couleur(ws.cell(i["ligne"], 1)) in marques or not code_de(i)]
    lignes_t = {i["ligne"] for i in a_traiter}
    modeles = [i for i in infos if i["ligne"] not in lignes_t]
    nb_f = lambda r: sum(txt(ws.cell(r, c).value).startswith("=") for c in range(1, ws.max_column + 1))
    tpl = max((m["ligne"] for m in modeles), key=nb_f, default=None)
    cols_f = [c for c in range(1, ws.max_column + 1) if tpl and txt(ws.cell(tpl, c).value).startswith("=")]
    index, utilises, par_ia = construire_index(sub, modeles), {code_de(m) for m in modeles}, set()
    types = {**apprendre_types(a_traiter, 0.8), **apprendre_types(modeles, 0.5)}  # lignes validees d'abord
    print(f"  Formats de detail appris : {sum(1 for t in types.values() if t)} sous-familles ; "
          f"{sum(1 for t in types.values() if not t)} a codes propres au produit.")
    print(f"  {len(modeles)} lignes modeles, {len(a_traiter)} lignes a traiter.")
    stats, debut = Counter(), time.time()

    for n, i in enumerate(a_traiter):
        r, ancien = i["ligne"], code_de(i) or "aucun"
        afficher_ligne(ws, r, sens)
        K = M = N = None
        try:
            classement = classer(i, index)
            cle, sure_sf, famille = choisir_sous_famille(i, classement, sub, modeles)
            if cle:
                K, M = cle.split("-", 1)
                if cle not in types:
                    types[cle] = type_par_ia(cle, sub, i, types)
                    par_ia.add(cle)
                N, probleme = construire_detail(i, cle, types[cle], utilises)
                statut = f"A VERIFIER : {probleme}" if probleme else "A VERIFIER : sous-famille incertaine" if not sure_sf \
                    else "A VERIFIER : type de detail choisi par l'IA" if cle in par_ia else "OK"
                if N:
                    utilises.add(f"{cle}-{N}")
                if statut == "OK" and i["four"].upper() in ("", "X"):
                    statut = "OK (sans designation fournisseur)"
            else:
                K, statut = famille, "A VERIFIER : sous-famille introuvable"
        except Exception as e:  # toute erreur : on marque la ligne et on continue
            print(f"     {ROUGE}erreur : {e}{FIN}")
            statut = "A VERIFIER : erreur IA"
        if tpl and cols_f and not all(txt(ws.cell(r, c).value).startswith("=") for c in cols_f):
            reparer_ligne(ws, r, tpl, cols_f)
        ws.cell(r, col["K"]).value, ws.cell(r, col["M"]).value, ws.cell(r, col["N"]).value = K, M, N or None
        ws.cell(r, col["statut"]).value = statut
        ws.cell(r, col["N"]).comment = Comment(f"Ancien code : {ancien}", "Codification IA")
        nouveau = f"{K}-{M}-{N}" if K and M and N else ""
        stats["ok" if statut.startswith("OK") else "verif"] += 1
        reste = (time.time() - debut) / (n + 1) * (len(a_traiter) - n - 1)
        print(f"    => {(VERT if statut.startswith('OK') else ROUGE)}{nouveau or '(pas de code)'}  [{statut}]{FIN}  "
              f"(reste ~{reste / 60:.1f} min)", flush=True)
        rapport.append([chemin.name, r, i["num"], i["nim"], ancien, nouveau, statut])

    # Collisions : doublons du code complet sur TOUT le fichier (+ couleur "collision" de la legende)
    groupes, teinte = {}, next(iter(roles["collision"]), None)
    for i in infos:
        ws.cell(i["ligne"], col["collision"]).value = None
        code = code_de(ligne_info(ws, i["ligne"], col))
        if code:
            groupes.setdefault(code, []).append(i["ligne"])
    nb_coll = 0
    for lignes in groupes.values():
        if len(lignes) > 1:
            nb_coll += len(lignes)
            for r in lignes:
                ws.cell(r, col["collision"]).value = f"OUI (lignes {', '.join(map(str, lignes[:8]))}{'...' if len(lignes) > 8 else ''})"
                if teinte:
                    ws.cell(r, col["collision"]).fill = PatternFill("solid", fgColor="FF" + teinte)
    sortie = DOSSIER_OUTPUT / f"{chemin.stem}_codifie_{horodatage}.xlsx"
    wb.save(sortie)
    print(f"\nResume : {len(a_traiter)} traitees | {stats['ok']} OK | {stats['verif']} a verifier | "
          f"{nb_coll} lignes en collision | {(time.time() - debut) / 60:.1f} min\nFichier : {sortie}\n")


def main():
    os.system("")  # active les couleurs ANSI sous Windows
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    DOSSIER_INPUT.mkdir(exist_ok=True)
    DOSSIER_OUTPUT.mkdir(exist_ok=True)
    fichiers = sorted(p for p in DOSSIER_INPUT.glob("*.xlsx") if not p.name.startswith("~$"))
    if not fichiers:
        sys.exit(f"Le dossier '{DOSSIER_INPUT.name}' ne contient aucun fichier .xlsx : depose ton fichier Excel dedans.")
    verifier_ollama()
    maintenant, rapport = datetime.now(), []
    for f in fichiers:
        print(f"\n=== {f.name} ===")
        traiter_fichier(f, maintenant.strftime("%Y-%m-%d_%Hh%M"), rapport)
    chemin_csv = DOSSIER_OUTPUT / f"rapport_{maintenant.strftime('%Y-%m-%d')}.csv"
    with open(chemin_csv, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["fichier", "ligne", "N°", "designation NIMRODA", "ancien code", "nouveau code", "statut"])
        w.writerows(rapport)
    print(f"Rapport : {chemin_csv}")


if __name__ == "__main__":
    main()
